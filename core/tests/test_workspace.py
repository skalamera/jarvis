"""Uploaded-file artifacts and codebase tools: parsing, versioned edits, undo, and the safety rails."""
import importlib
import io
import json

import pytest


@pytest.fixture()
def mods(tmp_path, monkeypatch):
    monkeypatch.setenv("JARVIS_STATE_DIR", str(tmp_path / "state"))
    monkeypatch.setenv("JARVIS_PROJECTS_DIR", str(tmp_path / "Projects"))
    monkeypatch.setenv("HOME", str(tmp_path))
    from jarvis_google import artifacts, code, store
    importlib.reload(store)
    importlib.reload(artifacts)
    importlib.reload(code)
    return artifacts, code, store


# ------------------------------------------------------------------ artifacts
def test_csv_parse_stats_and_versioned_edit(mods):
    A, _, store = mods
    m = A.save_upload("sales.csv", b"region,q1,q2\nEast,100,120\nWest,90,80\n", "text/csv")
    assert m["kind"] == "csv"
    b = A.file_open(m["id"])
    sheet = b["sheets"][0]
    assert sheet["header"] == ["region", "q1", "q2"]
    q1 = next(c for c in sheet["numeric_columns"] if c["name"] == "q1")
    assert q1["sum"] == 190 and q1["avg"] == 95
    r = A.sheet_edit(m["id"], [{"cell": "B2", "value": "150"}], append_rows=[["North", 5, 6]])
    assert r["saved_version"] == 2
    assert A.file_read(m["id"])["content"] == "region,q1,q2\nEast,150,120\nWest,90,80\nNorth,5,6\n"
    A.file_revert(m["id"])
    assert A.file_read(m["id"])["content"].startswith("region,q1,q2\nEast,100,")
    # every write is a card in the feed
    tools = [f["tool"] for f in store.feed_since(0)]
    assert tools == ["file_open", "sheet_edit", "file_revert"]


def test_xlsx_formulas_and_cell_types(mods):
    A, _, _ = mods
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.append(["a", "b", "sum"])
    wb.active.append([1, 2, "=A2+B2"])
    buf = io.BytesIO()
    wb.save(buf)
    m = A.save_upload("t.xlsx", buf.getvalue())
    v = A.view(A._meta(m["id"]))
    s = v["sheets"][0]
    assert s["grid"][1] == ["1", "2", "=A2+B2"] and s["formulas"] == {"1,2": "=A2+B2"}
    A.sheet_edit(m["id"], [{"cell": "A2", "value": "5"}, {"cell": "D1", "value": "=SUM(A2:B2)"}])
    ws = openpyxl.load_workbook(A.current_path(A._meta(m["id"]))).active
    assert ws["A2"].value == 5 and ws["D1"].value == "=SUM(A2:B2)"  # number stays a number, formula stays a formula


def test_text_edit_requires_exact_unique_match(mods):
    A, _, _ = mods
    m = A.save_upload("a.py", b"x = 1\ny = 1\n")
    with pytest.raises(ValueError, match="not found"):
        A.file_edit(m["id"], "z = 1", "z = 2")
    with pytest.raises(ValueError, match="matches 2"):
        A.file_edit(m["id"], "= 1", "= 2")
    assert A.file_edit(m["id"], "= 1", "= 2", replace_all=True)["replacements"] == 2


def test_image_ops(mods):
    A, _, _ = mods
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (200, 100), "blue").save(buf, "PNG")
    m = A.save_upload("p.png", buf.getvalue())
    r = A.image_edit(m["id"], [{"op": "rotate", "degrees": 90}, {"op": "crop", "left": 0, "top": 0, "right": 1, "bottom": 0.5}])
    assert (r["width"], r["height"]) == (100, 100)
    with pytest.raises(ValueError):
        A.image_edit(m["id"], [{"op": "explode"}])


def test_wrong_kind_and_bad_ids(mods):
    A, _, _ = mods
    m = A.save_upload("p.png", b"\x89PNG\r\n\x1a\n")
    with pytest.raises(ValueError):
        A.file_write(m["id"], "text")
    for bad in ("../../etc", "art_zzz", ""):
        with pytest.raises(ValueError):
            A.file_open(bad)


def test_upload_name_is_sanitised(mods):
    A, _, _ = mods
    m = A.save_upload("../../evil/../x.csv", b"a\n1\n")
    assert m["filename"] == "x.csv"
    assert A.current_path(m).parent.parent == A.ART_DIR


def test_export_never_overwrites(mods, tmp_path):
    A, _, _ = mods
    m = A.save_upload("r.txt", b"hi")
    a = A.export(m["id"])["path"]
    b = A.export(m["id"])["path"]
    assert a != b and b.endswith("r (2).txt")


# ------------------------------------------------------------------ code
def _project(tmp_path):
    root = tmp_path / "Projects" / "demo"
    (root / "pkg").mkdir(parents=True)
    (root / "web").mkdir()
    (root / "pkg" / "__init__.py").write_text("")
    (root / "pkg" / "core.py").write_text("def run():\n    return 1\n")
    (root / "main.py").write_text("from pkg.core import run\n\nprint(run())\n")
    (root / "web" / "util.ts").write_text("export function add(a: number, b: number) { return a + b }\n")
    (root / "web" / "index.ts").write_text("import { add } from './util'\nconsole.log(add(1, 2))\n")
    (root / ".env").write_text("SECRET=1\n")
    (root / "node_modules").mkdir()
    (root / "node_modules" / "junk.js").write_text("x\n" * 1000)
    return root


def test_code_map_modules_edges_and_ignores(mods, tmp_path):
    _, C, store = mods
    _project(tmp_path)
    b = C.code_map("demo")
    names = {m["name"] for m in b["modules"]}
    assert {"pkg", "web", "(root)"} <= names
    assert "node_modules" not in json.dumps(b)
    assert "(root) -> pkg (1)" in b["module_dependencies"]
    assert "main.py" in b["entry_points"]
    card = store.feed_since(0)[-1]["result"]
    files = {f["path"] for m in card["modules"] for f in m["files"]}
    assert "web/util.ts" in files and ".env" not in files
    util = next(f for m in card["modules"] for f in m["files"] if f["path"] == "web/util.ts")
    assert util["imported_by"] == 1 and "function add" in util["symbols"]


def test_code_annotate_persists(mods, tmp_path):
    _, C, store = mods
    _project(tmp_path)
    C.code_annotate("demo", "A demo.", {"pkg": "Core logic"}, ["main -> pkg.core.run"])
    card = store.feed_since(0)[-1]["result"]
    assert card["notes"]["summary"] == "A demo." and card["notes"]["modules"]["pkg"] == "Core logic"


def test_code_edit_diff_and_undo(mods, tmp_path):
    _, C, _ = mods
    root = _project(tmp_path)
    r = C.code_edit("demo", "pkg/core.py", "return 1", "return 2", note="bump")
    assert (root / "pkg/core.py").read_text() == "def run():\n    return 2\n"
    assert r["lines_added"] == 1 and r["lines_removed"] == 1
    C.undo_change(r["change_id"])
    assert (root / "pkg/core.py").read_text() == "def run():\n    return 1\n"
    with pytest.raises(ValueError, match="already undone"):
        C.undo_change(r["change_id"])


def test_undo_refuses_when_file_changed_since(mods, tmp_path):
    _, C, _ = mods
    root = _project(tmp_path)
    r = C.code_write("demo", "new.py", "a = 1\n")
    (root / "new.py").write_text("a = 1\n# his own edit\n")
    with pytest.raises(ValueError, match="changed since"):
        C.undo_change(r["change_id"])
    assert (root / "new.py").exists()


def test_code_safety_rails(mods, tmp_path):
    _, C, _ = mods
    _project(tmp_path)
    for bad in (".env", "../outside.py", ".git/config", "keys/server.pem", "/etc/hosts"):
        with pytest.raises(PermissionError):
            C.code_write("demo", bad, "x")
    with pytest.raises(PermissionError):
        C.code_read("demo", ".env")
    with pytest.raises(PermissionError):
        C.resolve_root(str(tmp_path))  # the home folder itself


def test_click_only_ops_are_not_mcp_tools():
    import inspect
    from jarvis_google import server
    src = inspect.getsource(server)
    assert "AR.export" not in src and "CODE.undo_change" not in src


def test_sheet_stats_skip_totals_row():
    from jarvis_google.artifacts import sheet_stats
    g = [["Account", "Size"], ["A", "10"], ["B", "30"], ["Total", "40"]]
    st = sheet_stats(g)
    assert st[0]["sum"] == 40 and st[0]["count"] == 2 and st[0]["max"] == 30
