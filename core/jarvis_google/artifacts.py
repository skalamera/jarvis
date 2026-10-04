"""Uploaded files (images, docs, spreadsheets, code) as versioned JARVIS artifacts.

Uploads are COPIES kept under the JARVIS state dir, so analysing and editing them never touches the original
on disk. Every modification writes a new version (v1, v2, ...) and the previous one stays, so any edit (by the
model or by a click) can be undone. Writing back OUT of the sandbox (Save to Downloads) is click-only.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import mimetypes
import re
import shutil
import time
import uuid
from pathlib import Path
from typing import Any

from . import store

ART_DIR = store.STATE_DIR / "artifacts"
MAX_UPLOAD = 50 * 1024 * 1024
GRID_ROWS, GRID_COLS = 2000, 60       # what the card shows / edits
TEXT_CAP = 400_000                     # characters shown in a text card
READ_CHUNK = 20_000                    # characters per file_read for the model

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tif", ".tiff"}
CSV_EXT = {".csv", ".tsv", ".tab"}
XLSX_EXT = {".xlsx", ".xlsm"}
CODE_EXT = {".py", ".ts", ".tsx", ".js", ".jsx", ".mjs", ".cjs", ".rs", ".go", ".java", ".kt", ".swift", ".c", ".h",
            ".cc", ".cpp", ".hpp", ".cs", ".rb", ".php", ".sh", ".zsh", ".bash", ".sql", ".html", ".css", ".scss",
            ".json", ".yaml", ".yml", ".toml", ".ini", ".xml", ".vue", ".svelte", ".lua", ".r", ".dart", ".scala"}
# iPhone (.mov/.mp4, HEVC/HDR), Android (.mp4/.3gp/.webm/.mkv), desktop (.avi/.wmv/.flv/.mpg/.mts/...)
VIDEO_EXT = {".mp4", ".mov", ".webm", ".m4v", ".3gp", ".3g2", ".mkv", ".avi", ".wmv", ".flv", ".mpg", ".mpeg",
             ".m2ts", ".mts", ".ts", ".ogv", ".hevc", ".qt", ".asf", ".vob", ".divx", ".f4v"}
MAX_VIDEO_UPLOAD = 4 * 1024 * 1024 * 1024
VIDEO_MIME = {".mov": "video/quicktime", ".qt": "video/quicktime", ".mkv": "video/x-matroska", ".3gp": "video/3gpp",
              ".3g2": "video/3gpp2", ".avi": "video/x-msvideo", ".wmv": "video/x-ms-wmv", ".flv": "video/x-flv",
              ".mts": "video/mp2t", ".m2ts": "video/mp2t", ".ts": "video/mp2t", ".m4v": "video/mp4"}
TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".svg"}
TEXT_KINDS = {"code", "text", "csv"}
_ID = re.compile(r"^art_[0-9a-f]{10}$")
_CELL = re.compile(r"^([A-Za-z]{1,3})(\d{1,7})$")


# ------------------------------------------------------------------ storage
def _clean_name(name: str) -> str:
    base = Path(str(name or "")).name
    base = re.sub(r"[^\w.\- ()]+", "_", base).strip(" .") or "untitled"
    return base[:120]


def kind_of(filename: str, head: bytes = b"") -> str:
    ext = Path(filename).suffix.lower()
    if ext in IMAGE_EXT:
        return "image"
    if ext in VIDEO_EXT:
        return "video"
    if ext in CSV_EXT:
        return "csv"
    if ext in XLSX_EXT:
        return "xlsx"
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext in CODE_EXT:
        return "code"
    if ext in TEXT_EXT:
        return "text"
    if head and b"\x00" not in head[:8192]:
        try:
            head[:8192].decode("utf-8")
            return "text"
        except UnicodeDecodeError:
            pass
    return "binary"


def _dir(aid: str) -> Path:
    if not _ID.match(str(aid or "")):
        raise ValueError(f"not an artifact id: {aid!r}")
    d = ART_DIR / aid
    if not (d / "meta.json").exists():
        raise FileNotFoundError(f"no such file: {aid}")
    return d


def _meta(aid: str) -> dict:
    return json.loads((_dir(aid) / "meta.json").read_text())


def _save_meta(m: dict) -> None:
    (ART_DIR / m["id"] / "meta.json").write_text(json.dumps(m, indent=1, default=str))


def current_path(m: dict) -> Path:
    return ART_DIR / m["id"] / m["versions"][-1]["file"]


def raw_path(aid: str, version: int | None = None) -> tuple[Path, str]:
    m = _meta(aid)
    v = m["versions"][-1] if not version else next((x for x in m["versions"] if x["n"] == version), m["versions"][-1])
    return ART_DIR / m["id"] / v["file"], m["mime"]


def _new_version(m: dict, data: bytes, source: str, note: str) -> dict:
    n = len(m["versions"]) + 1
    fname = f"v{n}{m['ext']}"
    (ART_DIR / m["id"] / fname).write_bytes(data)
    m["versions"].append({"n": n, "file": fname, "ts": time.time(), "source": source, "note": note[:200],
                          "size": len(data)})
    m["size"], m["updated_at"] = len(data), time.time()
    _save_meta(m)
    if source != "upload":
        store.audit({"kind": "artifact_version", "artifact": m["id"], "filename": m["filename"], "version": n,
                     "source": source, "note": note[:200]})
    return m


def is_video_name(filename: str) -> bool:
    return Path(str(filename or "")).suffix.lower() in VIDEO_EXT


def save_upload_file(filename: str, src: Path, mime: str | None = None) -> dict:
    """Big uploads (videos) streamed to a temp file first: moved into the sandbox instead of held in memory."""
    if src.stat().st_size > MAX_VIDEO_UPLOAD:
        raise ValueError(f"file is larger than {MAX_VIDEO_UPLOAD // (1024 ** 3)} GB")
    name = _clean_name(filename)
    ext = Path(name).suffix.lower()
    aid = "art_" + uuid.uuid4().hex[:10]
    (ART_DIR / aid).mkdir(parents=True, exist_ok=True)
    shutil.move(str(src), ART_DIR / aid / f"v1{ext}")
    size = (ART_DIR / aid / f"v1{ext}").stat().st_size
    m = {"id": aid, "filename": name, "ext": ext, "kind": kind_of(name),
         "mime": VIDEO_MIME.get(ext) or (mime if mime and mime != "application/octet-stream" else None)
         or mimetypes.guess_type(name)[0] or "application/octet-stream",
         "created_at": time.time(), "updated_at": time.time(), "size": size,
         "versions": [{"n": 1, "file": f"v1{ext}", "ts": time.time(), "source": "upload", "note": "original", "size": size}]}
    _save_meta(m)
    return m


def save_upload(filename: str, data: bytes, mime: str | None = None, aid: str = "", source: str = "upload",
                note: str = "original") -> dict:
    """A new workspace file. aid / source / note are for generated files (the id is reserved up front so the progress
    card and the finished file share a display)."""
    if len(data) > MAX_UPLOAD and source == "upload":
        raise ValueError(f"file is larger than {MAX_UPLOAD // (1024 * 1024)} MB")
    name = _clean_name(filename)
    aid = aid if _ID.match(aid or "") else "art_" + uuid.uuid4().hex[:10]
    (ART_DIR / aid).mkdir(parents=True, exist_ok=True)
    m = {"id": aid, "filename": name, "ext": Path(name).suffix.lower(), "kind": kind_of(name, data),
         "mime": mime or mimetypes.guess_type(name)[0] or "application/octet-stream",
         "created_at": time.time(), "updated_at": time.time(), "size": len(data), "versions": []}
    return _new_version(m, data, source, note)


def list_artifacts(limit: int = 30) -> list[dict]:
    out = []
    if ART_DIR.exists():
        for d in ART_DIR.iterdir():
            try:
                m = json.loads((d / "meta.json").read_text())
            except Exception:
                continue
            out.append({"id": m["id"], "filename": m["filename"], "kind": m["kind"], "size": m["size"],
                        "version": len(m["versions"]), "updated_at": m["updated_at"]})
    return sorted(out, key=lambda x: x["updated_at"], reverse=True)[:limit]


# ------------------------------------------------------------------ extraction
def _decode(data: bytes) -> str:
    if data[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return data.decode("utf-16")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _cell_str(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.10g}"
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return v.isoformat()
    return str(v)


def _num(s: str) -> float | None:
    t = str(s).strip().replace(",", "").replace("$", "").replace("%", "")
    if not t:
        return None
    try:
        return float(t)
    except ValueError:
        return None


def sheet_stats(grid: list[list[str]]) -> list[dict]:
    """Numeric summary per column (header = first row). A column counts as numeric when >= 60% of filled cells are."""
    if len(grid) < 2:
        return []
    header, body = grid[0], grid[1:]
    # a trailing "Total" / "Grand total" / "Sum" row is a summary, not data: counting it would double every sum
    while body and any(re.fullmatch(r"\s*(grand\s+)?(totals?|sum|subtotal)\s*:?\s*", str(c), re.I) for c in body[-1][:3]):
        body = body[:-1]
    out = []
    for ci, name in enumerate(header):
        vals = [r[ci] for r in body if ci < len(r) and str(r[ci]).strip()]
        nums = [n for n in (_num(v) for v in vals) if n is not None]
        if vals and len(nums) >= max(1, 0.6 * len(vals)):
            out.append({"col": ci, "name": name or f"Column {ci + 1}", "count": len(nums), "sum": round(sum(nums), 4),
                        "avg": round(sum(nums) / len(nums), 4), "min": min(nums), "max": max(nums)})
    return out


def _delimiter(m: dict, text: str) -> str:
    if m["ext"] in (".tsv", ".tab"):
        return "\t"
    try:
        return csv.Sniffer().sniff(text[:8192], delimiters=",;\t|").delimiter
    except csv.Error:
        return ","


def _csv_grid(m: dict, text: str) -> tuple[list[list[str]], str]:
    delim = _delimiter(m, text)
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)], delim


def _xlsx_sheets(path: Path) -> list[dict]:
    import openpyxl
    vals = openpyxl.load_workbook(path, data_only=True)
    forms = openpyxl.load_workbook(path, data_only=False)
    sheets = []
    for ws in vals.worksheets:
        fws = forms[ws.title]
        grid: list[list[str]] = []
        formulas: dict[str, str] = {}
        for ri, row in enumerate(ws.iter_rows(max_row=min(ws.max_row, GRID_ROWS), max_col=min(ws.max_column, GRID_COLS))):
            out_row = []
            for ci, cell in enumerate(row):
                f = fws.cell(row=ri + 1, column=ci + 1).value
                if isinstance(f, str) and f.startswith("="):
                    formulas[f"{ri},{ci}"] = f
                # a formula never computed by Excel has no cached value: show the formula itself
                out_row.append(_cell_str(cell.value) if cell.value is not None else (f if isinstance(f, str) else ""))
            grid.append(out_row)
        while grid and not any(c.strip() for c in grid[-1]):
            grid.pop()
        sheets.append({"name": ws.title, "grid": grid, "formulas": formulas, "total_rows": ws.max_row,
                       "total_cols": ws.max_column, "stats": sheet_stats(grid)})
    return sheets


def _pdf_text(path: Path) -> tuple[str, int]:
    from pypdf import PdfReader
    r = PdfReader(str(path))
    parts = []
    for i, p in enumerate(r.pages):
        try:
            parts.append(f"--- page {i + 1} ---\n{p.extract_text() or ''}")
        except Exception as e:
            parts.append(f"--- page {i + 1} --- (unreadable: {e})")
    return "\n".join(parts), len(r.pages)


def _docx_text(path: Path) -> tuple[str, int]:
    import docx
    d = docx.Document(str(path))
    paras = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            paras.append(" | ".join(c.text for c in row.cells))
    return "\n".join(paras), len(d.tables)


def full_text(m: dict) -> str:
    """The file as text, for reading (sheets as CSV, PDFs / Word as extracted text)."""
    p, k = current_path(m), m["kind"]
    if k in TEXT_KINDS:
        return _decode(p.read_bytes())
    if k == "xlsx":
        out = []
        for s in _xlsx_sheets(p):
            buf = io.StringIO()
            csv.writer(buf, lineterminator="\n").writerows(s["grid"])
            out.append(f"### sheet: {s['name']} ({s['total_rows']} rows x {s['total_cols']} cols)\n{buf.getvalue()}")
        return "\n".join(out)
    if k == "pdf":
        return _pdf_text(p)[0]
    if k == "docx":
        return _docx_text(p)[0]
    return ""


def view(m: dict) -> dict:
    p, k = current_path(m), m["kind"]
    try:
        if k == "image":
            from PIL import Image
            with Image.open(p) as im:
                return {"type": "image", "width": im.width, "height": im.height, "format": im.format, "mode": im.mode}
        if k == "video":
            return {"type": "video", "size": m["size"]}
        if k == "csv":
            grid, delim = _csv_grid(m, _decode(p.read_bytes()))
            return {"type": "sheet", "delimiter": delim, "sheets": [{
                "name": m["filename"], "grid": [r[:GRID_COLS] for r in grid[:GRID_ROWS]], "formulas": {},
                "total_rows": len(grid), "total_cols": max((len(r) for r in grid), default=0),
                "stats": sheet_stats(grid)}]}
        if k == "xlsx":
            return {"type": "sheet", "sheets": _xlsx_sheets(p)}
        if k in ("code", "text"):
            text = _decode(p.read_bytes())
            return {"type": "text", "text": text[:TEXT_CAP], "truncated": len(text) > TEXT_CAP,
                    "lines": text.count("\n") + 1, "language": m["ext"].lstrip(".") or "text"}
        if k == "pdf":
            text, pages = _pdf_text(p)
            return {"type": "document", "text": text[:TEXT_CAP], "truncated": len(text) > TEXT_CAP, "pages": pages}
        if k == "docx":
            text, tables = _docx_text(p)
            return {"type": "document", "text": text[:TEXT_CAP], "truncated": len(text) > TEXT_CAP, "tables": tables}
    except Exception as e:
        return {"type": "error", "error": f"{type(e).__name__}: {e}"[:300]}
    return {"type": "binary", "size": m["size"]}


def card_data(m: dict) -> dict:
    return {"key": f"artifact:{m['id']}",
            "artifact": {"id": m["id"], "filename": m["filename"], "kind": m["kind"], "ext": m["ext"], "mime": m["mime"],
                         "size": m["size"], "version": len(m["versions"]), "created_at": m["created_at"],
                         "updated_at": m["updated_at"],
                         "versions": [{k: v[k] for k in ("n", "ts", "source", "note")} for v in m["versions"]][-12:]},
            "view": view(m)}


def brief(m: dict, v: dict | None = None) -> dict:
    """Compact description for the model (the card gets the full data)."""
    v = v if v is not None else view(m)
    out: dict[str, Any] = {"id": m["id"], "filename": m["filename"], "kind": m["kind"], "size_bytes": m["size"],
                           "version": len(m["versions"])}
    t = v.get("type")
    if t == "sheet":
        out["sheets"] = [{"name": s["name"], "rows": s["total_rows"], "cols": s["total_cols"],
                          "header": s["grid"][0] if s["grid"] else [], "first_rows": s["grid"][1:8],
                          "numeric_columns": s["stats"][:12], "formula_count": len(s.get("formulas") or {})}
                         for s in v["sheets"][:8]]
    elif t in ("text", "document"):
        out["lines"] = v.get("lines")
        out["pages"] = v.get("pages")
        out["excerpt"] = v.get("text", "")[:3000]
        if len(v.get("text", "")) > 3000:
            out["more"] = "file_read for the rest"
    elif t == "image":
        out.update({"width": v["width"], "height": v["height"], "format": v["format"], "path": str(current_path(m)),
                    "how_to_see_it": "Call your vision_analyze tool with this path to look at the image."})
    elif t == "error":
        out["error"] = v["error"]
    return {k: x for k, x in out.items() if x is not None}


def _show(tool: str, args: dict, m: dict, record: bool = True) -> dict:
    data = card_data(m)
    if record:
        store.record_result(tool, None, args, data)
    return data


# ------------------------------------------------------------------ reads (model + clicks)
def file_open(artifact_id: str, show: bool = True) -> dict:
    m = _meta(artifact_id)
    data = _show("file_open", {"artifact_id": artifact_id}, m, record=show)
    return brief(m, data["view"])


def file_list(limit: int = 30) -> dict:
    return {"files": list_artifacts(limit)}


def file_read(artifact_id: str, offset: int = 0, limit: int = READ_CHUNK) -> dict:
    m = _meta(artifact_id)
    if m["kind"] == "image":
        return brief(m)
    if m["kind"] == "binary":
        return {"id": m["id"], "filename": m["filename"], "error": "binary file; no text to read"}
    text = full_text(m)
    offset, limit = max(0, int(offset)), max(1000, min(int(limit), 60_000))
    out = {"id": m["id"], "filename": m["filename"], "kind": m["kind"], "version": len(m["versions"]),
           "total_chars": len(text), "offset": offset, "content": text[offset:offset + limit]}
    if offset + limit < len(text):
        out["next_offset"] = offset + limit
    return out


# ------------------------------------------------------------------ writes (all reversible: new version each time)
def _write_text(m: dict, text: str, source: str, note: str) -> dict:
    if m["kind"] not in TEXT_KINDS:
        hint = {"xlsx": "sheet_edit", "image": "image_edit"}.get(m["kind"], "file_create for an edited copy")
        raise ValueError(f"{m['filename']} is a {m['kind']} file; use {hint}")
    return _new_version(m, text.encode("utf-8"), source, note)


def file_write(artifact_id: str, content: str, note: str = "", source: str = "model") -> dict:
    m = _write_text(_meta(artifact_id), content, source, note or "rewritten")
    _show("file_write", {"artifact_id": artifact_id}, m)
    return {"id": m["id"], "filename": m["filename"], "saved_version": len(m["versions"])}


def file_edit(artifact_id: str, old: str, new: str, replace_all: bool = False, note: str = "",
              source: str = "model") -> dict:
    m = _meta(artifact_id)
    text = full_text(m)
    n = text.count(old) if old else 0
    if n == 0:
        raise ValueError("old text not found (it must match exactly, including whitespace); file_read it first")
    if n > 1 and not replace_all:
        raise ValueError(f"old text matches {n} places; add surrounding context or set replace_all")
    done = n if replace_all else 1
    m = _write_text(m, text.replace(old, new) if replace_all else text.replace(old, new, 1), source,
                    note or f"edited {done} place(s)")
    _show("file_edit", {"artifact_id": artifact_id}, m)
    return {"id": m["id"], "filename": m["filename"], "saved_version": len(m["versions"]), "replacements": done}


def _col_index(letters: str) -> int:
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


def _target(e: dict) -> tuple[int, int]:
    if e.get("cell"):
        mm = _CELL.match(str(e["cell"]).strip())
        if not mm:
            raise ValueError(f"bad cell reference {e['cell']!r} (use A1 style)")
        return int(mm.group(2)) - 1, _col_index(mm.group(1))
    r, c = int(e["row"]), int(e["col"])
    if r < 0 or c < 0:
        raise ValueError("row/col must be >= 0")
    return r, c


def _coerce(v: Any) -> Any:
    if isinstance(v, str):
        s = v.strip()
        if s.startswith("="):
            return s
        if re.fullmatch(r"-?\d+", s):
            return int(s)
        if re.fullmatch(r"-?\d*\.\d+(e-?\d+)?|-?\d+e-?\d+", s, re.I):
            return float(s)
    return v


def sheet_edit(artifact_id: str, edits: list[dict] | None = None, sheet: str = "", append_rows: list[list] | None = None,
               delete_rows: list[int] | None = None, note: str = "", source: str = "model") -> dict:
    """edits: [{cell:"B3", value}] or [{row, col, value}] (0-based grid row/col, row 0 = header)."""
    m = _meta(artifact_id)
    edits, append_rows = edits or [], append_rows or []
    deletes = sorted({int(r) for r in (delete_rows or [])}, reverse=True)
    if not (edits or append_rows or deletes):
        raise ValueError("nothing to change")
    p = current_path(m)
    if m["kind"] == "csv":
        grid, delim = _csv_grid(m, _decode(p.read_bytes()))
        for e in edits:
            r, c = _target(e)
            while len(grid) <= r:
                grid.append([])
            while len(grid[r]) <= c:
                grid[r].append("")
            grid[r][c] = _cell_str(e.get("value"))
        for r in deletes:
            if 0 <= r < len(grid):
                del grid[r]
        grid.extend([[_cell_str(x) for x in row] for row in append_rows])
        buf = io.StringIO()
        csv.writer(buf, delimiter=delim, lineterminator="\n").writerows(grid)
        m = _new_version(m, buf.getvalue().encode("utf-8"), source, note or f"{len(edits)} cell edit(s)")
    elif m["kind"] == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(p, keep_vba=m["ext"] == ".xlsm")
        if sheet and sheet not in wb.sheetnames:
            raise ValueError(f"no sheet {sheet!r}; sheets: {wb.sheetnames}")
        ws = wb[sheet] if sheet else wb.active
        for e in edits:
            r, c = _target(e)
            ws.cell(row=r + 1, column=c + 1, value=_coerce(e.get("value")))
        for r in deletes:
            ws.delete_rows(r + 1)
        for row in append_rows:
            ws.append([_coerce(x) for x in row])
        buf = io.BytesIO()
        wb.save(buf)
        m = _new_version(m, buf.getvalue(), source, note or f"{len(edits)} cell edit(s) on {ws.title}")
    else:
        raise ValueError(f"{m['filename']} is not a spreadsheet")
    _show("sheet_edit", {"artifact_id": artifact_id}, m)
    return {"id": m["id"], "filename": m["filename"], "saved_version": len(m["versions"]), "cells_changed": len(edits),
            "rows_appended": len(append_rows), "rows_deleted": len(deletes)}


IMAGE_OPS = {"rotate", "flip", "grayscale", "resize", "crop", "brightness", "contrast", "saturation", "sharpen", "blur",
             "annotate", "invert"}


def _apply_image_ops(im, ops: list[dict]):
    from PIL import Image, ImageDraw, ImageEnhance, ImageFilter, ImageFont, ImageOps
    for op in ops:
        kind = str(op.get("op", "")).lower()
        if kind not in IMAGE_OPS:
            raise ValueError(f"unknown image op {kind!r}; use one of {sorted(IMAGE_OPS)}")
        if kind == "rotate":
            im = im.rotate(-float(op.get("degrees", 90)), expand=True)  # positive = clockwise
        elif kind == "flip":
            im = ImageOps.mirror(im) if str(op.get("direction", "horizontal")).startswith("h") else ImageOps.flip(im)
        elif kind == "grayscale":
            im = ImageOps.grayscale(im).convert("RGB")
        elif kind == "invert":
            im = ImageOps.invert(im.convert("RGB"))
        elif kind == "resize":
            w, h = op.get("width"), op.get("height")
            if not w and not h:
                raise ValueError("resize needs width and/or height")
            w = int(w or im.width * int(h) / im.height)
            h = int(h or im.height * int(w) / im.width)
            im = im.resize((max(1, w), max(1, h)), Image.Resampling.LANCZOS)
        elif kind == "crop":
            box = [float(op.get(k, d)) for k, d in (("left", 0), ("top", 0), ("right", 1), ("bottom", 1))]
            if all(0 <= b <= 1 for b in box):  # fractions of the image
                box = [box[0] * im.width, box[1] * im.height, box[2] * im.width, box[3] * im.height]
            l, t, r, b = (int(x) for x in box)
            if r <= l or b <= t:
                raise ValueError("crop box is empty")
            im = im.crop((l, t, r, b))
        elif kind in ("brightness", "contrast", "saturation"):
            enh = {"brightness": ImageEnhance.Brightness, "contrast": ImageEnhance.Contrast,
                   "saturation": ImageEnhance.Color}[kind]
            im = enh(im if im.mode in ("RGB", "RGBA") else im.convert("RGB")).enhance(float(op.get("factor", 1.2)))
        elif kind == "sharpen":
            im = im.filter(ImageFilter.SHARPEN)
        elif kind == "blur":
            im = im.filter(ImageFilter.GaussianBlur(float(op.get("radius", 2))))
        elif kind == "annotate":
            im = im if im.mode in ("RGB", "RGBA") else im.convert("RGB")
            d = ImageDraw.Draw(im)
            size = int(op.get("size", max(14, im.width // 30)))
            try:
                font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", size)
            except Exception:
                font = ImageFont.load_default()
            x, y = float(op.get("x", 0.03)), float(op.get("y", 0.03))
            x, y = (x * im.width if x <= 1 else x), (y * im.height if y <= 1 else y)
            d.text((x, y), str(op.get("text", ""))[:200], fill=str(op.get("color", "#ff3b30")), font=font,
                   stroke_width=max(1, size // 12), stroke_fill="black")
    return im


def image_edit(artifact_id: str, ops: list[dict], note: str = "", source: str = "model") -> dict:
    from PIL import Image
    m = _meta(artifact_id)
    if m["kind"] != "image":
        raise ValueError(f"{m['filename']} is not an image")
    if not ops:
        raise ValueError("no ops")
    with Image.open(current_path(m)) as src:
        fmt = src.format or "PNG"
        im = _apply_image_ops(src.copy(), ops)
    if fmt.upper() in ("JPEG", "JPG") and im.mode not in ("RGB", "L"):
        im = im.convert("RGB")
    buf = io.BytesIO()
    im.save(buf, format=fmt)
    m = _new_version(m, buf.getvalue(), source, note or ", ".join(str(o.get("op")) for o in ops))
    _show("image_edit", {"artifact_id": artifact_id}, m)
    return {"id": m["id"], "filename": m["filename"], "saved_version": len(m["versions"]), "width": im.width,
            "height": im.height, "path": str(current_path(m))}


def file_create(filename: str, content: str = "", rows: list[list] | None = None, note: str = "") -> dict:
    """A NEW file in the JARVIS workspace (report, CSV, script, cleaned copy...). `rows` builds an .xlsx/.csv grid."""
    name = _clean_name(filename)
    ext = Path(name).suffix.lower()
    if rows is not None and ext in XLSX_EXT:
        import openpyxl
        wb = openpyxl.Workbook()
        for r in rows:
            wb.active.append([_coerce(x) for x in r])
        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
    elif rows is not None:
        buf = io.StringIO()
        csv.writer(buf, lineterminator="\n").writerows(rows)
        data = buf.getvalue().encode("utf-8")
    else:
        data = (content or "").encode("utf-8")
    m = save_upload(name, data)
    m["versions"][0]["source"], m["versions"][0]["note"] = "model", note[:200] or "created by JARVIS"
    _save_meta(m)
    store.audit({"kind": "artifact_create", "artifact": m["id"], "filename": name})
    _show("file_create", {"filename": name}, m)
    return brief(m)


def file_revert(artifact_id: str, source: str = "model") -> dict:
    m = _meta(artifact_id)
    if len(m["versions"]) < 2:
        raise ValueError("already at the original version")
    prev = m["versions"][-2]
    m = _new_version(m, (ART_DIR / m["id"] / prev["file"]).read_bytes(), source, f"undo: back to v{prev['n']}")
    _show("file_revert", {"artifact_id": artifact_id}, m)
    return {"id": m["id"], "filename": m["filename"], "saved_version": len(m["versions"]), "restored_from": prev["n"]}


def export(artifact_id: str) -> dict:
    """Click-only: copy the current version to ~/Downloads (never overwrites; adds ' (2)' etc.)."""
    m = _meta(artifact_id)
    dl = Path.home() / "Downloads"
    dl.mkdir(exist_ok=True)
    stem, ext = Path(m["filename"]).stem, m["ext"]
    dest, i = dl / m["filename"], 2
    while dest.exists():
        dest, i = dl / f"{stem} ({i}){ext}", i + 1
    shutil.copyfile(current_path(m), dest)
    store.audit({"kind": "artifact_export", "artifact": m["id"], "dest": str(dest), "source": "artifact_click"})
    return {"path": str(dest), "text": f"Saved to Downloads as {dest.name}."}


# ------------------------------------------------------------------ HUD click ops (Core allow-list)
def click_get(artifact_id: str) -> dict:
    return card_data(_meta(artifact_id))


def click_save_text(artifact_id: str, content: str) -> dict:
    _write_text(_meta(artifact_id), content, "artifact_click", "edited in the HUD")
    return {**click_get(artifact_id), "text": "Saved."}


def click_sheet_set(artifact_id: str, edits: list[dict] | None = None, sheet: str = "",
                    append_rows: list[list] | None = None, delete_rows: list[int] | None = None) -> dict:
    sheet_edit(artifact_id, edits, sheet, append_rows, delete_rows, "edited in the HUD", "artifact_click")
    return {**click_get(artifact_id), "text": "Saved."}


def click_image_op(artifact_id: str, ops: list[dict]) -> dict:
    image_edit(artifact_id, ops, "", "artifact_click")
    return {**click_get(artifact_id), "text": "Applied."}


def click_revert(artifact_id: str) -> dict:
    file_revert(artifact_id, "artifact_click")
    return {**click_get(artifact_id), "text": "Undone."}


def click_discard(artifact_id: str) -> dict:
    """Click-only: he dismissed a freshly generated image/video in the center hologram. Only JARVIS-generated,
    never-edited files qualify (his uploads and edited files are refused). Moved to a trash dir, not hard-deleted."""
    m = _meta(artifact_id)
    vs = m.get("versions") or []
    if len(vs) != 1 or vs[0].get("source") != "generated":
        raise ValueError("only an unedited JARVIS-generated file can be discarded")
    trash = store.STATE_DIR / "artifacts_trash"
    trash.mkdir(parents=True, exist_ok=True)
    shutil.move(str(_dir(artifact_id)), trash / f"{artifact_id}-{int(time.time())}")
    store.audit({"kind": "artifact_discard", "artifact": artifact_id, "filename": m["filename"], "source": "artifact_click"})
    return {"text": f"Discarded {m['filename']}.", "artifact_id": artifact_id}


CLICK_OPS = {"artifact_get": click_get, "artifact_discard": click_discard, "artifact_save_text": click_save_text, "artifact_sheet_set": click_sheet_set,
             "artifact_image_op": click_image_op, "artifact_revert": click_revert, "artifact_export": export}
