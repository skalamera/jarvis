"""Stephen's car (2003 Mercedes-Benz CL600, C215): a private garage knowledge base built from his Google Drive folder
(receipts, owner's manual, diagrams, guides) plus general model knowledge, and the tools JARVIS uses to be an expert
on it: profile / specs / service history / mods / maintenance due, document search, photo & video diagnosis, upgrade
plans, and logging new work.

Everything lives in STATE_DIR/garage (outside the repo; contains his VIN and paperwork):
  files/<drive_id>.<ext>   downloaded copies          docs.json     per-file extraction (Gemini)
  text/<drive_id>.txt      full text for search        profile.json  aggregated vehicle profile
  plans.json               upgrade plans               log.json      work he reports by voice
Config: JARVIS_CAR_DRIVE_FOLDER (Drive folder id) in env / ~/.hermes/.env; Drive account "personal".
"""
from __future__ import annotations

import base64
import datetime as dt
import io
import json
import math
import os
import re
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from typing import Any

import httpx

from . import store

G = store.STATE_DIR / "garage"
FILES, TEXT = G / "files", G / "text"
DOCS, PROFILE, PLANS, LOG = G / "docs.json", G / "profile.json", G / "plans.json", G / "log.json"
FACTS = G / "owner_facts.json"  # things HE has told JARVIS about the car; authoritative over documents
SEED_FACTS = ["The ABC (Active Body Control) hydraulic suspension has been removed and replaced with Strutmasters coil "
              "struts (conversion kit MC14FM): the car no longer has ABC, so ABC fluid, accumulators, pump, valve blocks, "
              "level sensors and lowering links no longer apply.",
              "The car has no front or rear sway bars installed; he reports it drives fine without them."]
ACCOUNT = "personal"
MODEL = os.environ.get("JARVIS_GARAGE_MODEL", "gemini-3.8-flash")
API = "https://generativelanguage.googleapis.com/v1beta"
CAR = "2003 Mercedes-Benz CL600 (C215 chassis, M275 5.5L twin-turbo V12, 5-speed 722.6 automatic)"
_lock = threading.Lock()


# ------------------------------------------------------------------ config / io
def _env(name: str) -> str:
    v = os.environ.get(name, "").strip()
    if v:
        return v
    f = Path(os.environ.get("HERMES_HOME", Path.home() / ".hermes")) / ".env"
    try:
        for line in f.read_text().splitlines():
            m = re.match(rf"\s*(?:export\s+)?{name}\s*=\s*(.*)", line)
            if m and m[1].strip():
                return m[1].strip().strip('"').strip("'")
    except OSError:
        pass
    return ""


def _folder() -> str:
    f = _env("JARVIS_CAR_DRIVE_FOLDER")
    if not f:
        raise RuntimeError("No car folder configured: set JARVIS_CAR_DRIVE_FOLDER in ~/.hermes/.env.")
    return f


def _load(p: Path, default: Any) -> Any:
    try:
        return json.loads(p.read_text())
    except Exception:
        return default


def _save(p: Path, data: Any) -> None:
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".tmp")
    tmp.write_text(json.dumps(data, indent=1, default=str))
    tmp.replace(p)
    os.chmod(p, 0o600)


def _gemini(parts: list[dict], schema_hint: str = "", json_out: bool = True, timeout: float = 180) -> Any:
    key = _env("GEMINI_API_KEY")
    if not key:
        raise RuntimeError("GEMINI_API_KEY is not configured")
    cfg: dict[str, Any] = {"temperature": 0.2}
    if json_out:
        cfg["responseMimeType"] = "application/json"
    for attempt in range(3):
        r = httpx.post(f"{API}/models/{MODEL}:generateContent", headers={"x-goog-api-key": key}, timeout=timeout,
                       json={"contents": [{"parts": parts}], "generationConfig": cfg})
        if r.status_code in (429, 500, 503) and attempt < 2:
            time.sleep(4 * (attempt + 1))
            continue
        if r.status_code != 200:
            raise RuntimeError(f"Gemini {r.status_code}: {r.text[:300]}")
        cand = (r.json().get("candidates") or [{}])[0]
        text = "".join(p.get("text", "") for p in (cand.get("content") or {}).get("parts", []))
        if not json_out:
            return text
        text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text.strip())
        try:
            return json.loads(text)
        except json.JSONDecodeError:
            m = re.search(r"[\[{].*[\]}]", text, re.S)
            if m:
                return json.loads(m.group(0))
            raise
    raise RuntimeError("Gemini unavailable")


def _upload_file(data: bytes, mime: str, name: str) -> str:
    """Gemini File API (large videos / PDFs). Returns the file uri once ACTIVE."""
    key = _env("GEMINI_API_KEY")
    start = httpx.post("https://generativelanguage.googleapis.com/upload/v1beta/files", timeout=60,
                       headers={"x-goog-api-key": key, "X-Goog-Upload-Protocol": "resumable",
                                "X-Goog-Upload-Command": "start", "X-Goog-Upload-Header-Content-Length": str(len(data)),
                                "X-Goog-Upload-Header-Content-Type": mime, "Content-Type": "application/json"},
                       json={"file": {"display_name": name[:100]}})
    url = start.headers.get("x-goog-upload-url")
    if not url:
        raise RuntimeError(f"Gemini upload failed ({start.status_code})")
    up = httpx.post(url, timeout=300, content=data, headers={"X-Goog-Upload-Offset": "0",
                                                             "X-Goog-Upload-Command": "upload, finalize"})
    f = up.json()["file"]
    for _ in range(60):
        if f.get("state") == "ACTIVE":
            return f["uri"]
        if f.get("state") == "FAILED":
            raise RuntimeError("Gemini couldn't process that file")
        time.sleep(2)
        f = httpx.get(f"{API}/{f['name']}", headers={"x-goog-api-key": key}, timeout=30).json()
    raise RuntimeError("Gemini file processing timed out")


def _media_part(data: bytes, mime: str, name: str) -> dict:
    if len(data) <= 18 * 1024 * 1024:
        return {"inlineData": {"mimeType": mime, "data": base64.b64encode(data).decode()}}
    return {"fileData": {"mimeType": mime, "fileUri": _upload_file(data, mime, name)}}


# ------------------------------------------------------------------ Drive sync
def _drive():
    from .accounts import service
    return service("drive", ACCOUNT)


def _walk(d, fid: str, path: str = "") -> list[dict]:
    out, tok = [], None
    while True:
        r = d.files().list(q=f"'{fid}' in parents and trashed=false", pageSize=200, pageToken=tok,
                           fields="nextPageToken,files(id,name,mimeType,size,modifiedTime,webViewLink,thumbnailLink)",
                           supportsAllDrives=True, includeItemsFromAllDrives=True).execute()
        for f in r.get("files", []):
            if f["mimeType"] == "application/vnd.google-apps.folder":
                out += _walk(d, f["id"], f"{path}/{f['name'].strip()}" if path else f["name"].strip())
            else:
                out.append({**f, "folder": path})
        tok = r.get("nextPageToken")
        if not tok:
            return out


_EXPORT = {"application/vnd.google-apps.document": ("application/pdf", ".pdf"),
           "application/vnd.google-apps.spreadsheet": ("text/csv", ".csv")}


def _download(d, f: dict) -> tuple[Path, str] | None:
    mt = f["mimeType"]
    if mt in _EXPORT:
        emime, ext = _EXPORT[mt]
        data = d.files().export(fileId=f["id"], mimeType=emime).execute()
        mime = emime
    elif mt.startswith("application/vnd.google-apps"):
        return None  # shortcuts etc.
    else:
        data = d.files().get_media(fileId=f["id"], supportsAllDrives=True).execute()
        mime = mt
        ext = {"application/pdf": ".pdf", "image/jpeg": ".jpg", "image/png": ".png", "image/heic": ".heic",
               "video/mp4": ".mp4", "video/quicktime": ".mov",
               "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx"}.get(mt, "")
    FILES.mkdir(parents=True, exist_ok=True)
    p = FILES / f"{f['id']}{ext}"
    p.write_bytes(data)
    return p, mime


def _local_text(p: Path) -> tuple[str, int]:
    if p.suffix == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(str(p))
        out = []
        for i, pg in enumerate(r.pages):
            try:
                out.append(f"[page {i + 1}]\n{pg.extract_text() or ''}")
            except Exception:
                out.append(f"[page {i + 1}]")
        return "\n".join(out), len(r.pages)
    if p.suffix == ".docx":
        from docx import Document
        doc = Document(str(p))
        lines = [x.text for x in doc.paragraphs if x.text.strip()]
        for t in doc.tables:
            for row in t.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(lines), 0
    if p.suffix == ".csv":
        return p.read_text(errors="replace"), 0
    return "", 0


EXTRACT_PROMPT = """You are cataloguing one document from the owner's records for his {car}.
File: "{name}" (folder: {folder}).
Return JSON:
{{"doc_type": "receipt|invoice|estimate|owners_manual|diagram|spec_sheet|guide|registration|title|bill_of_sale|insurance|photo|history_report|warranty|other",
 "title": "short human title",
 "date": "YYYY-MM-DD or null (the service / purchase date on the document)",
 "mileage": integer odometer reading on the document or null,
 "vendor": "shop / seller name or null",
 "total_cost": number in USD or null,
 "work": [{{"item": "what was done / bought", "category": "maintenance|repair|upgrade|parts|inspection|other",
           "system": "engine|transmission|suspension|brakes|electrical|cooling|fuel|exhaust|steering|interior|exterior|wheels_tires|hvac|other",
           "parts": [{{"name": "...", "part_number": "... or null", "brand": "... or null", "qty": number or null, "price": number or null}}]}}],
 "facts": ["concrete reusable facts: specs, capacities, torque values, fluid types, fuse assignments, intervals, VIN, colors, codes"],
 "summary": "2-3 sentences",
 "transcript": "{transcript_rule}"}}
Use only what's in the document. Null when unknown. Part numbers exactly as printed."""


def _extract(f: dict, p: Path, mime: str) -> dict:
    text, pages = _local_text(p)
    long_doc = pages > 12 or len(text) > 60000
    scanned = not long_doc and len(text.strip()) < 250 * max(1, pages)  # image-only PDF / photo: transcribe it
    rule = ("verbatim readable text of the document, including every table row (fuse number, amperage, circuit), "
            "label and callout on diagrams, and every receipt line item and total" if scanned
            else "empty string (the full text is extracted separately)")
    prompt = EXTRACT_PROMPT.format(car=CAR, name=f["name"], folder=f.get("folder") or "/", transcript_rule=rule)
    parts: list[dict] = [{"text": prompt}]
    if long_doc:
        parts.append({"text": "DOCUMENT TEXT (first part):\n" + text[:120000]})
    elif mime.startswith(("image/", "video/")) or mime == "application/pdf":
        parts.append(_media_part(p.read_bytes(), mime, f["name"]))
        if text.strip():
            parts.append({"text": "Extracted text layer:\n" + text[:40000]})
    else:
        parts.append({"text": text[:120000] or "(no text)"})
    try:
        x = _gemini(parts)
    except Exception as e:
        x = {"doc_type": "other", "title": f["name"], "summary": f"(couldn't read: {str(e)[:120]})", "work": [], "facts": []}
    if isinstance(x, list):
        x = x[0] if x else {}
    full = text
    if scanned:
        full = _transcribe(p, mime, f["name"]) or x.get("transcript") or text
    TEXT.mkdir(parents=True, exist_ok=True)
    (TEXT / f"{f['id']}.txt").write_text(full)
    x.pop("transcript", None)
    return {**x, "id": f["id"], "name": f["name"], "folder": f.get("folder") or "", "mime": mime, "pages": pages,
            "file": p.name, "link": f.get("webViewLink"), "modified": f.get("modifiedTime"), "chars": len(full)}


def reextract(match: str = "", scanned_only: bool = True) -> list[str]:
    """Re-read existing files (e.g. after improving extraction) without re-downloading."""
    with _lock:
        docs = {x["id"]: x for x in _load(DOCS, [])}
        redo = [x for x in docs.values() if (not match or re.search(match, x["name"] + " " + (x.get("folder") or ""), re.I))
                and (not scanned_only or x.get("chars", 0) < 250 * max(1, x.get("pages") or 1))]
        f = lambda x: _extract({"id": x["id"], "name": x["name"], "folder": x.get("folder"), "webViewLink": x.get("link"),
                                "modifiedTime": x.get("modified")}, FILES / x["file"], x["mime"])
        with ThreadPoolExecutor(6) as ex:
            for r in ex.map(f, redo):
                docs[r["id"]] = r
        _save(DOCS, sorted(docs.values(), key=lambda x: (x.get("folder") or "", x.get("name") or "")))
        _chunks_cache.clear()
    return [x["name"] for x in redo]


TRANSCRIBE_PROMPT = """Transcribe ALL readable text in this document from a car owner's records, page by page.
Start each page with a line "[page N]". Reproduce tables as rows with " | " between cells, keeping every row
(e.g. fuse number | amperage | protected circuit). Include labels and callouts on diagrams. No commentary."""


def _transcribe(p: Path, mime: str, name: str) -> str:
    try:
        return _gemini([{"text": TRANSCRIBE_PROMPT}, _media_part(p.read_bytes(), mime, name)], json_out=False,
                       timeout=300).strip()
    except Exception:
        return ""


def sync(force: bool = False, rebuild: bool = False, max_workers: int = 6) -> dict:
    """Pull new / changed files from his Drive car folder, extract them, rebuild the profile."""
    with _lock:
        d = _drive()
        files = _walk(d, _folder())
        docs = {x["id"]: x for x in _load(DOCS, [])}
        todo = [f for f in files if force or f["id"] not in docs or docs[f["id"]].get("modified") != f.get("modifiedTime")]

        def work(f: dict) -> dict | None:
            got = _download(_drive(), f)
            return _extract(f, *got) if got else None
        with ThreadPoolExecutor(max_workers) as ex:
            for r in ex.map(work, todo):
                if r:
                    docs[r["id"]] = r
        live = {f["id"] for f in files}
        removed = [k for k in docs if k not in live]
        for k in removed:
            docs.pop(k, None)
        _save(DOCS, sorted(docs.values(), key=lambda x: (x.get("folder") or "", x.get("name") or "")))
        _chunks_cache.clear()
        prof = build_profile() if (rebuild or todo or removed or not PROFILE.exists()) else _load(PROFILE, {})
    return {"files": len(files), "new_or_changed": len(todo), "removed": len(removed),
            "service_entries": len(prof.get("service_history") or []), "mods": len(prof.get("mods") or [])}


# ------------------------------------------------------------------ profile (aggregate)
PROFILE_PROMPT = """You are building the definitive profile of ONE specific car for its owner's assistant:
{car}. Below: structured extracts of every document in his records (receipts, manual, diagrams), plus work he logged.
Combine them with your expert knowledge of the C215 CL600 / M275 engine. Return JSON:
{{
 "vehicle": {{"year": 2003, "make": "Mercedes-Benz", "model": "CL600", "chassis": "C215", "engine": "...", "transmission": "...",
             "vin": "... or null", "color": "... or null", "interior": "... or null", "purchased": "YYYY-MM-DD or null",
             "purchase_price": number or null, "current_mileage": latest odometer seen (int) or null,
             "mileage_as_of": "YYYY-MM-DD or null", "state": "registration state or null", "nickname": null}},
 "configuration_notes": ["how THIS car differs from stock (e.g. suspension conversion, tune, removed parts) - critical for advice"],
 "specs": [{{"group": "Engine|Drivetrain|Performance|Dimensions|Capacities & Fluids|Wheels & Tires|Electrical|Suspension|Brakes",
             "label": "...", "value": "...", "source": "doc|stock"}}],
 "service_history": [{{"date": "YYYY-MM-DD", "mileage": int or null, "vendor": "...", "title": "...", "category": "maintenance|repair|upgrade|inspection",
                       "system": "...", "items": ["..."], "parts": [{{"name": "...", "part_number": "..."}}], "cost": number or null, "doc_ids": ["..."]}}],
 "mods": [{{"name": "...", "category": "performance|suspension|lighting|exterior|interior|audio|wheels_tires|other", "date": "YYYY-MM-DD or null",
           "vendor": "... or null", "cost": number or null, "details": "...", "doc_ids": ["..."]}}],
 "maintenance": [{{"item": "...", "interval_miles": int or null, "interval_months": int or null, "last_date": "YYYY-MM-DD or null",
                   "last_mileage": int or null, "spec": "fluid / part spec (e.g. MB 229.5 5W-40, 8.5 qt)", "notes": "...", "source": "doc|stock"}}],
 "known_issues": [{{"title": "...", "system": "...", "severity": "low|medium|high", "applies": "yes|no|resolved|watch",
                    "why": "how it relates to THIS car's history/config", "symptoms": "...", "fix": "..."}}],
 "documents": [{{"id": "...", "title": "...", "category": "Receipts|Manuals|Diagrams & Specs|Paperwork|Photos|Other", "date": "YYYY-MM-DD or null"}}],
 "totals": {{"documented_spend": number, "since": "YYYY-MM-DD"}}
}}
Rules: service_history = one entry per real service visit / purchase (merge duplicate photos of the same receipt), newest first.
Never invent dates, costs, mileage or part numbers; null when unknown. Mark spec / interval source "doc" only if a document states it.
known_issues: the well-known C215 / M275 failure points (ABC, ignition coils & harness, turbo vacuum/boost lines, transmission
conductor plate, intercooler pump, SAM modules, COMAND, door/window regulators, etc.), each marked against this car's history.

DOCUMENT EXTRACTS:
{docs}

OWNER-LOGGED WORK:
{log}

OWNER-STATED FACTS (authoritative: they override anything older in the documents; configuration_notes must reflect them,
and maintenance / known_issues must not list items for systems the car no longer has, or mark them applies "no"):
{facts}"""


def build_profile() -> dict:
    docs = _load(DOCS, [])
    slim = [{k: x.get(k) for k in ("id", "name", "folder", "doc_type", "title", "date", "mileage", "vendor", "total_cost",
                                   "work", "facts", "summary")} for x in docs]
    prof = _gemini([{"text": PROFILE_PROMPT.format(car=CAR, docs=json.dumps(slim, default=str)[:400000],
                                                   log=json.dumps(_load(LOG, []), default=str),
                                                   facts=json.dumps(owner_facts()))}], timeout=300)
    links = {x["id"]: x for x in docs}
    _ground_costs(prof, links)
    for e in (prof.get("service_history") or []) + (prof.get("mods") or []):
        e["docs"] = [{"id": i, "title": links[i].get("title") or links[i]["name"], "link": links[i].get("link")}
                     for i in e.get("doc_ids") or [] if i in links]
    for dd in prof.get("documents") or []:
        x = links.get(dd.get("id")) or {}
        dd.update({"link": x.get("link"), "name": x.get("name"), "mime": x.get("mime"), "summary": x.get("summary"),
                   "folder": x.get("folder")})
    have = {dd.get("id") for dd in prof.get("documents") or []}
    for x in docs:  # never drop a document from the index
        if x["id"] not in have:
            (prof.setdefault("documents", [])).append({"id": x["id"], "title": x.get("title") or x["name"],
                                                       "category": "Other", "date": x.get("date"), "link": x.get("link"),
                                                       "name": x["name"], "mime": x.get("mime"), "summary": x.get("summary"),
                                                       "folder": x.get("folder")})
    prof["hero"] = _pick_hero(docs)
    prof["built_at"] = time.time()
    _save(PROFILE, prof)
    return prof


def owner_facts() -> list[str]:
    f = _load(FACTS, None)
    if f is None:
        _save(FACTS, SEED_FACTS)
        return list(SEED_FACTS)
    return f


def add_fact(fact: str) -> dict:
    facts = owner_facts()
    if fact not in facts:
        facts.append(fact)
        _save(FACTS, facts)
    return {"saved": fact, "facts": facts,
            "note": "Saved. It applies on the next profile rebuild (car_sync rebuild=true) and to every answer now."}


_PURCHASE = re.compile(r"\b(vehicle purchase|bill of sale|purchase of vehicle)\b", re.I)


_PAID = {"receipt", "invoice"}
_WORDS = re.compile(r"[a-z0-9]{3,}")


def _bag(*xs: Any) -> set[str]:
    return set(_WORDS.findall(json.dumps(xs, default=str).lower())) - _STOP


def _ground_costs(prof: dict, links: dict) -> None:
    """Costs, estimate flags and spend totals come from the per-document extracts (read straight off each receipt),
    never from the aggregate model pass. Each receipt counts once, against the entry its contents match best;
    owner-compiled summaries / insurance statements never set a cost (they'd double count)."""
    entries = [e for e in prof.get("service_history") or [] if e.get("source") != "logged"]
    owner: dict[str, int] = {}
    for i, e in enumerate(entries):
        eb = _bag(e.get("title"), e.get("items"))
        for d in e.get("doc_ids") or []:
            x = links.get(d)
            if not x:
                continue
            score = len(eb & _bag(x.get("title"), x.get("summary"), [w.get("item") for w in x.get("work") or []]))
            if d not in owner or score > owner[d][1]:
                owner[d] = (i, score)
    total, since = 0.0, None
    for i, e in enumerate(entries):
        ds = [links[d] for d, (j, _) in owner.items() if j == i]
        paid = [x for x in ds if x.get("doc_type") in _PAID and isinstance(x.get("total_cost"), (int, float))]
        est = [x for x in ds if x.get("doc_type") == "estimate" and isinstance(x.get("total_cost"), (int, float))]
        use = paid or est
        e["cost"] = round(sum({round(x["total_cost"], 2) for x in use}), 2) if use else None  # same receipt twice = once
        e["estimate"] = bool(est) and not paid
        e["purchase"] = bool(_PURCHASE.search(e.get("title") or ""))
        if e["cost"] and not e["estimate"] and not e["purchase"]:
            total += e["cost"]
            since = min(since or "9999", e.get("date") or "9999")
    for e in prof.get("service_history") or []:
        if e.get("source") == "logged" and e.get("cost"):
            total += e["cost"]
    for m in prof.get("mods") or []:
        ds = [links[i] for i in m.get("doc_ids") or [] if i in links]
        paid = [x for x in ds if x.get("doc_type") in _PAID and isinstance(x.get("total_cost"), (int, float))]
        m["cost"] = round(sum({round(x["total_cost"], 2) for x in paid}), 2) if paid else (
            m.get("cost") if m.get("source") == "logged" else None)
    prof["totals"] = {"documented_spend": round(total, 2), "since": since if since != "9999" else None,
                      "note": "paid invoices & receipts only (excludes the purchase price and estimates / insurance quotes)"}


HERO = G / "hero.png"
HERO_PROMPT = ("Photorealistic three-quarter front studio shot of a 2003 Mercedes-Benz CL600 coupe (C215, pillarless "
               "two-door, Brilliant Silver Metallic, stock 18-inch wheels), slightly lowered, on a dark reflective floor "
               "with soft cyan rim light, cinematic, no text, no people.")


def _pick_hero(docs: list[dict]) -> str | None:
    """A real photo of the car from his files (doc_type photo) wins; else a one-time studio render."""
    for x in docs:
        if (x.get("mime") or "").startswith("image/") and x.get("doc_type") == "photo":
            return x["id"]
    if not HERO.exists():
        try:
            key = _env("GEMINI_API_KEY")
            from .generate import IMAGE_MODELS
            r = httpx.post(f"{API}/models/{IMAGE_MODELS['fast']}:generateContent", headers={"x-goog-api-key": key},
                           timeout=180, json={"contents": [{"parts": [{"text": HERO_PROMPT}]}],
                                              "generationConfig": {"responseModalities": ["IMAGE"],
                                                                   "imageConfig": {"aspectRatio": "4:3"}}})
            for part in ((r.json().get("candidates") or [{}])[0].get("content") or {}).get("parts", []):
                if "inlineData" in part:
                    HERO.write_bytes(base64.b64decode(part["inlineData"]["data"]))
                    break
        except Exception:
            return None
    return "hero" if HERO.exists() else None


def profile() -> dict:
    p = _load(PROFILE, None)
    if p is None:
        raise RuntimeError("The car profile hasn't been built yet: run car_sync first.")
    return _with_status(p)


def _with_status(p: dict) -> dict:
    """Maintenance due status from the last RECORDED odometer reading (never extrapolated) + today's date."""
    v = p.get("vehicle") or {}
    today = dt.date.today()
    est = v.get("current_mileage")
    for m in p.get("maintenance") or []:
        due, pct = [], None
        if m.get("interval_miles") and m.get("last_mileage") and est:
            left = m["last_mileage"] + m["interval_miles"] - est
            due.append(left / max(1, m["interval_miles"]))
            m["miles_left"] = int(left)
        if m.get("interval_months") and m.get("last_date"):
            try:
                last = dt.date.fromisoformat(m["last_date"])
                months = (today.year - last.year) * 12 + today.month - last.month
                left_m = m["interval_months"] - months
                due.append(left_m / max(1, m["interval_months"]))
                m["months_left"] = left_m
            except ValueError:
                pass
        if due:
            pct = min(due)
        m["status"] = ("unknown" if pct is None else "overdue" if pct < 0 else "due_soon" if pct < 0.15 else "ok")
        m["remaining_pct"] = None if pct is None else max(0.0, min(1.0, pct))
    p["estimated_mileage"] = est
    return p


# ------------------------------------------------------------------ search over his documents
_chunks_cache: list[dict] = []


def _chunks() -> list[dict]:
    if _chunks_cache:
        return _chunks_cache
    for x in _load(DOCS, []):
        t = ""
        try:
            t = (TEXT / f"{x['id']}.txt").read_text()
        except OSError:
            pass
        head = f"{x.get('title') or x['name']} | {x.get('summary') or ''} | " + " ; ".join(x.get("facts") or [])
        _chunks_cache.append({"doc": x, "page": None, "text": head})
        for m in re.finditer(r"(?:\[page (\d+)\]\n)?((?:(?!\[page \d+\]).){1,1600})", t, re.S):
            body = m.group(2).strip()
            if len(body) > 40:
                _chunks_cache.append({"doc": x, "page": int(m.group(1)) if m.group(1) else None, "text": body})
    return _chunks_cache


_STOP = set("the a an and or of to for in on my is it how what when which do does i car cl600 mercedes with at be".split())


def _terms(q: str) -> list[str]:
    return [w for w in re.findall(r"[a-z0-9][a-z0-9\-\.]+", q.lower()) if w not in _STOP]


def search(query: str, limit: int = 6, record: bool = True) -> dict:
    """Keyword search (BM25-ish) across every page of his car documents."""
    ch = _chunks()
    if not ch:
        raise RuntimeError("No car documents indexed yet: run car_sync first.")
    terms = _terms(query)
    phrase = re.sub(r"[\s\-\.]", "", query.lower())
    digits = len(re.sub(r"\D", "", phrase)) >= 6  # part numbers: match ignoring spaces / dashes
    n = len(ch)
    df = {t: sum(1 for c in ch if t in c["text"].lower()) for t in terms}
    scored = []
    for c in ch:
        low = c["text"].lower()
        s = 0.0
        for t in terms:
            tf = low.count(t)
            if tf:
                s += math.log(1 + n / (1 + df[t])) * (tf * 2.2) / (tf + 1.2 * (0.25 + 0.75 * len(low) / 900))
        if len(phrase) >= 5 and phrase in re.sub(r"[\s\-\.]", "", low):
            s = s * 3 + (25 if digits else 8)  # exact phrase / part number
        if s > 0:
            scored.append((s * (0.6 if c["page"] is None and c["text"].count(" | ") >= 2 else 1.0), c))  # pages beat summaries
    scored.sort(key=lambda x: -x[0])
    hits, seen = [], set()
    for s, c in scored:
        k = (c["doc"]["id"], c["page"])
        if k in seen:
            continue
        seen.add(k)
        txt = c["text"]
        i = min((txt.lower().find(t) for t in terms if t in txt.lower()), default=0)
        if len(phrase) >= 5:  # centre the excerpt on an exact phrase / part-number match
            pm = re.search(r"[\s\-\.]*".join(map(re.escape, phrase)), txt, re.I)
            if pm:
                i = pm.start()
        snip = txt[max(0, i - 220): i + 650].strip()
        hits.append({"doc": c["doc"].get("title") or c["doc"]["name"], "doc_id": c["doc"]["id"],
                     "type": c["doc"].get("doc_type"), "page": c["page"], "link": c["doc"].get("link"),
                     "date": c["doc"].get("date"), "excerpt": snip, "score": round(s, 2)})
        if len(hits) >= limit:
            break
    res = {"query": query, "hits": hits}
    if record and hits:
        store.record_result("car_search", None, {"query": query}, {"key": f"garage:search:{query[:40]}",
                                                                   "kind": "garage_search", **res})
    return {"query": query, "hits": [{k: h[k] for k in ("doc", "page", "date", "excerpt")} for h in hits],
            "note": "Answer from these excerpts (cite the document); if they don't cover it, use expert knowledge and say so."}


# ------------------------------------------------------------------ display + compact model views
SECTIONS = ("overview", "specs", "service", "mods", "maintenance", "issues", "docs", "plans")


def car_profile(section: str = "overview", record: bool = True) -> dict:
    p = profile()
    sec = section if section in SECTIONS else "overview"
    plans = _load(PLANS, [])
    if record:
        store.record_result("car_profile", None, {"section": sec}, {
            "key": "garage:cl600", "kind": "garage", "tab": sec, "profile": p, "plans": plans,
            "hero_id": p.get("hero"), "updated": p.get("built_at")})
    v = p.get("vehicle") or {}
    hist = p.get("service_history") or []
    out: dict[str, Any] = {"vehicle": v, "last_recorded_mileage": p.get("estimated_mileage"),
                           "owner_facts": owner_facts(), "documented_spend": p.get("totals"),
                           "configuration_notes": p.get("configuration_notes"),
                           "last_service": hist[0] if hist else None, "mods": [m.get("name") for m in p.get("mods") or []],
                           "due": [{k: m.get(k) for k in ("item", "status", "miles_left", "months_left", "spec")}
                                   for m in p.get("maintenance") or [] if m.get("status") in ("overdue", "due_soon")]}
    if sec == "specs":
        out["specs"] = p.get("specs")
    elif sec == "service":
        out["service_history"] = [{k: e.get(k) for k in ("date", "mileage", "vendor", "title", "items", "cost")} for e in hist]
    elif sec == "mods":
        out["mods"] = p.get("mods")
    elif sec == "maintenance":
        out["maintenance"] = p.get("maintenance")
    elif sec == "issues":
        out["known_issues"] = p.get("known_issues")
    elif sec == "docs":
        out["documents"] = [{k: d.get(k) for k in ("title", "category", "date")} for d in p.get("documents") or []]
    elif sec == "plans":
        out["plans"] = plans
    out["shown"] = f"the {sec} tab of his car display is on screen"
    return out


# ------------------------------------------------------------------ diagnosis from photos / video
DIAG_PROMPT = """You are a master Mercedes-Benz technician specialising in the C215 CL600 (M275 twin-turbo V12).
Diagnose the problem shown for THIS specific car. Owner's description: "{symptoms}"
THIS CAR (use it: history, mods and configuration change the likely causes):
{context}
Look closely at the attached photo(s)/video (gauges, warning lights, leaks, colours of fluids, parts, sounds in video).
Return JSON:
{{"summary": "one-sentence read of what you see",
 "observations": ["specific things visible / audible in the media"],
 "urgency": "safe_to_drive|drive_carefully|stop_driving",
 "causes": [{{"cause": "...", "likelihood": 0-100, "why": "evidence + how this car's history makes it more/less likely",
             "system": "...", "fix": "...", "parts": [{{"name": "...", "part_number": "MB part number if confident, else null"}}],
             "diy": "easy|moderate|hard|shop", "est_cost": "rough USD range"}}],
 "checks": ["ordered next diagnostic steps he can do (codes to read, what to inspect / measure)"],
 "related_history": ["past work on this car that relates (with dates)"],
 "follow_up_question": "the one question whose answer would narrow it down most, or null"}}
Rank causes by likelihood (3-5). Be concrete; never invent part numbers."""


def _context_text(p: dict) -> str:
    v = p.get("vehicle") or {}
    hist = p.get("service_history") or []
    lines = [f"Last recorded odometer {v.get('current_mileage')} mi on {v.get('mileage_as_of')}",
             "Owner-stated facts (authoritative): " + " ".join(owner_facts()),
             "Configuration: " + "; ".join(p.get("configuration_notes") or []),
             "Mods: " + "; ".join(f"{m.get('name')} ({m.get('date') or '?'})" for m in p.get("mods") or []),
             "Service history:"] + [f"- {e.get('date')} @{e.get('mileage') or '?'} mi: {e.get('title')}: "
                                    f"{', '.join(e.get('items') or [])[:200]}" for e in hist[:40]]
    return "\n".join(lines)


def diagnose(symptoms: str = "", artifact_ids: list[str] | None = None) -> dict:
    from . import artifacts as AR
    p = profile()
    parts: list[dict] = [{"text": DIAG_PROMPT.format(symptoms=symptoms or "(none given; read the media)",
                                                     context=_context_text(p))}]
    media = []
    for aid in (artifact_ids or [])[:6]:
        m = AR._meta(aid)
        if m["kind"] not in ("image", "video"):
            continue
        data = AR.current_path(m).read_bytes()
        parts.append(_media_part(data, m["mime"] or ("video/mp4" if m["kind"] == "video" else "image/jpeg"), m["filename"]))
        media.append({"id": aid, "kind": m["kind"], "version": len(m["versions"]), "filename": m["filename"]})
    if not media and not symptoms:
        raise ValueError("Give me a photo / video (attach it) or describe the symptoms.")
    r = _gemini(parts, timeout=300)
    if isinstance(r, list):
        r = r[0] if r else {}
    r["causes"] = sorted(r.get("causes") or [], key=lambda c: -(c.get("likelihood") or 0))
    res = {"key": f"garage:diag:{uuid.uuid4().hex[:6]}", "kind": "garage_diagnosis", "symptoms": symptoms,
           "media": media, **r}
    store.record_result("car_diagnose", None, {"symptoms": symptoms[:120]}, res)
    return {k: r.get(k) for k in ("summary", "urgency", "observations", "checks", "follow_up_question")} | {
        "causes": [{k: c.get(k) for k in ("cause", "likelihood", "fix", "diy", "est_cost")} for c in r["causes"]],
        "shown": "the diagnosis display is on screen; give him the top cause, urgency and first check in 2-3 sentences"}


# ------------------------------------------------------------------ upgrade plans + logging work
def _cost(v: Any) -> tuple[float, float]:
    """'$1,800 - $2,600' / '1800' / 1800 -> (low, high)."""
    if isinstance(v, (int, float)):
        return float(v), float(v)
    nums = [float(x.replace(",", "")) * (1000 if k.lower() == "k" else 1)
            for x, k in re.findall(r"(\d[\d,]*(?:\.\d+)?)\s*([kK]?)", str(v or ""))]
    return (min(nums), max(nums)) if nums else (0.0, 0.0)


def _norm_stage(s: dict) -> dict:
    """Accept the shapes models actually send (parts: [str], est_cost ranges as strings) -> canonical stage."""
    raw = s.get("items") or s.get("parts") or []
    items = []
    for it in raw:
        it = {"part": it} if isinstance(it, str) else dict(it)
        it["part"] = it.get("part") or it.get("name") or it.get("item") or ""
        lo, hi = _cost(it.get("est_cost") if it.get("est_cost") is not None else it.get("price"))
        it["est_cost"], it["est_cost_high"] = (lo or None), (hi if hi != lo else None)
        items.append(it)
    lo = sum(i["est_cost"] or 0 for i in items)
    hi = sum(i.get("est_cost_high") or i["est_cost"] or 0 for i in items)
    if not lo and s.get("est_cost") is not None:  # one stage-level estimate
        lo, hi = _cost(s.get("est_cost"))
    return {"name": s.get("name") or s.get("title") or "Stage", "items": items,
            "notes": s.get("notes") or s.get("description") or "",
            "est_total": round(lo, 2), "est_total_high": round(hi, 2) if hi != lo else None}


def plan_save(title: str, goal: str, stages: list[dict], considerations: list[str] | None = None,
              plan_id: str = "") -> dict:
    """stages: [{name, items: [{part, brand, part_number, est_cost, labor_hours, notes}], notes}]"""
    plans = _load(PLANS, [])
    pid = plan_id or "plan_" + uuid.uuid4().hex[:8]
    stages = [_norm_stage(s) for s in stages or []]
    plan = {"id": pid, "title": title, "goal": goal, "stages": stages, "considerations": considerations or [],
            "total": round(sum(s["est_total"] for s in stages), 2),
            "total_high": round(sum(s.get("est_total_high") or s["est_total"] for s in stages), 2),
            "updated": time.time()}
    if plan["total_high"] == plan["total"]:
        plan["total_high"] = None
    plans = [x for x in plans if x["id"] != pid] + [plan]
    _save(PLANS, plans)
    store.record_result("car_plan", None, {"title": title}, {"key": f"garage:plan:{pid}", "kind": "garage_plan", **plan})
    return {"plan_id": pid, "total": plan["total"], "total_high": plan["total_high"], "stages": len(stages),
            "shown": "the build plan is on screen"}


def log_work(summary: str, date: str = "", mileage: int = 0, vendor: str = "", cost: float = 0,
             items: list[str] | None = None, kind: str = "maintenance") -> dict:
    """Record work he tells JARVIS about (no receipt in Drive yet). Folded into the profile."""
    log = _load(LOG, [])
    entry = {"date": date or dt.date.today().isoformat(), "mileage": int(mileage) or None, "vendor": vendor or None,
             "title": summary, "items": items or [summary], "cost": cost or None, "category": kind, "source": "logged"}
    log.append(entry)
    _save(LOG, log)
    p = _load(PROFILE, {})
    if p:
        p.setdefault("service_history", []).insert(0, {**entry, "docs": []})
        if kind == "upgrade":
            p.setdefault("mods", []).append({"name": summary, "date": entry["date"], "cost": cost or None,
                                             "vendor": vendor or None, "details": ", ".join(items or [])})
        if entry["mileage"] and entry["mileage"] > ((p.get("vehicle") or {}).get("current_mileage") or 0):
            p["vehicle"]["current_mileage"], p["vehicle"]["mileage_as_of"] = entry["mileage"], entry["date"]
        p["service_history"].sort(key=lambda e: e.get("date") or "", reverse=True)
        _save(PROFILE, p)
    return {"logged": entry, "note": "Saved to his car's history (and shown next time the service tab opens)."}


def file_path(doc_id: str) -> tuple[Path, str]:
    """A downloaded document, for the HUD (thumbnails / hero / opening receipts)."""
    if doc_id == "hero" and HERO.exists():
        return HERO, "image/png"
    for x in _load(DOCS, []):
        if x["id"] == doc_id:
            return FILES / x["file"], x.get("mime") or "application/octet-stream"
    raise FileNotFoundError(doc_id)
